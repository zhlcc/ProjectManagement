from flask import Flask, render_template, request, redirect, url_for, send_file
import psutil
import db_query
import datetime
from environment_rank import total_rank
import openpyxl

app = Flask(__name__)
app.config['SECRET_KEY'] = 'gsolvit'


@app.route('/')
def index():
    return render_template('index.html')


@app.route('/admin_login', methods=['GET', 'POST'])
def admin_login():
    if request.method == 'GET':
        return render_template('admin_login.html')
    else:
        admin_id = request.form.get('admin_id')
        password = request.form.get('password')
        sql = 'select* from admin where admin_id=%s and password=%s'
        result = db_query.db_query(sql, [admin_id, password])
        print(result)
        if len(result) != 0:
            if result[0]['password'] == password:
                return render_template('admin_index.html')
            else:
                error = '账号或密码错误'
        else:
            error = '账号或密码错误'
        return render_template('admin_login.html', error=error)


@app.route('/admin', methods=['GET', 'POST'])
def admin():
    if request.method == 'GET':
        return render_template('admin_index.html')
    else:
        table_name = request.form.get('table-name')
        print(table_name)
        sql = 'select* from ' + table_name
        data = db_query.db_query(sql)
        return render_template('admin_index.html', data=data, table_name=table_name)


@app.route('/add_item', methods=['POST'])
def add_item():
    table_name = request.form['table_name']
    form_data = request.form.to_dict()
    del form_data['table_name']
    print(table_name)
    print(form_data)
    columns = ', '.join(form_data.keys())
    placeholders = ', '.join(['%s'] * len(form_data))
    insert_sql = f"INSERT INTO {table_name} ({columns}) VALUES ({placeholders})"
    db_query.db_query(insert_sql, list(form_data.values()))
    sql = 'select* from ' + table_name
    data = db_query.db_query(sql)
    return render_template('admin_index.html', data=data, table_name=table_name)


@app.route('/delete_item', methods=['POST'])
def delete_item():
    table_name = request.form['table_name']
    row = request.form['row']
    print(table_name)
    print(row)
    sql_ = 'select* from ' + table_name
    data = db_query.db_query(sql_)
    d_item = data[int(row) - 1]
    print(d_item)
    where_clause = ' AND '.join([f"{key} = %s" for key in d_item])
    delete_sql = f"DELETE FROM {table_name} WHERE {where_clause}"
    print(delete_sql)
    values = list(d_item.values())
    print(values)
    db_query.db_query(delete_sql, values)
    data = db_query.db_query(sql_)
    return render_template('admin_index.html', data=data, table_name=table_name)


@app.route('/update_item', methods=['POST'])
def update_item():
    table_name = request.form['table_name']
    row = request.form['row']
    column = request.form['column']
    value = request.form['value']
    print(table_name)
    print(row)
    print(column)
    print(value)
    sql_ = 'select* from ' + table_name
    data = db_query.db_query(sql_)
    d_item = data[int(row) - 1]
    print(d_item)
    where_clause = ' AND '.join([f"{key} = %s" for key in d_item])
    update_sql = f"UPDATE {table_name} SET {column} = %s WHERE {where_clause}"
    print(update_sql)
    args = [value] + list(d_item.values())
    print(args)
    db_query.db_query(update_sql, args)
    data_ = db_query.db_query(sql_)
    return render_template('admin_index.html', data=data_, table_name=table_name)


@app.route('/user_login', methods=['GET', 'POST'])
def user_login():
    if request.method == 'GET':
        return render_template('user_login.html')
    else:
        user_id = request.form.get('user_id')
        password = request.form.get('password')
        sql = 'select* from user where user_id=%s and password=%s'
        result = db_query.db_query(sql, [user_id, password])
        print(result)
        if len(result) != 0:
            return redirect(url_for(('user')))
        else:
            error = '账号或密码错误'
            return render_template('user_login.html', error=error)


@app.route('/user', methods=['GET', 'POST'])
def user():
    sql = 'select* from 水质数据 where date=%s'
    result = db_query.db_query(sql, datetime.date.today())
    print(result)
    if request.method == 'GET':
        return render_template('user_index.html', water_quality=result)
    else:
        start_date_str = request.form.get('startDate')
        end_date_str = request.form.get('endDate')
        feature = request.form.get('feature')
        print(start_date_str, end_date_str, feature)
        start_date = datetime.datetime.strptime(start_date_str, '%Y-%m-%d').date()
        end_date = datetime.datetime.strptime(end_date_str, '%Y-%m-%d').date()
        sql = 'select * from 水质数据 where date BETWEEN %s AND %s'
        result1 = db_query.db_query(sql, [start_date, end_date])
        for data in result1:
            print(data)
        return render_template('user_index.html', water_quality=result, history_data=result1, feature=feature)


@app.route('/underwater_system', methods=['GET', 'POST'])
def underwater_system():
    sql = 'select* from 水质数据 where date=%s'
    result = db_query.db_query(sql, datetime.date.today())
    print(result)
    rank = total_rank(result[0])
    sql1 = 'select* from fish_num'
    nums = db_query.db_query(sql1)
    print(nums)
    sql2 = 'SELECT Species,COUNT(*) AS SpeciesNum FROM fish GROUP BY Species'
    species_num = db_query.db_query(sql2)
    print(species_num)
    species = [item['Species'] for item in species_num]
    if request.method == 'GET':
        return render_template('underwater_system.html', rank=rank, nums=nums, specie_num=len(species_num),
                               species=species, species_num=species_num)
    else:
        specie = request.form.get('specie')
        feature = request.form.get('feature')
        print(specie)
        print(feature)
        sql3 = 'SELECT * FROM fish WHERE species=%s'
        result = db_query.db_query(sql3, [specie])
        print(result)
        data = [item[feature] for item in result]
        print(data)
        return render_template('underwater_system.html', rank=rank, nums=nums, specie_num=len(species_num),
                               species=species, data=data, feature=feature, fish=specie, species_num=species_num)


@app.route('/data_center')
def data_center():
    sql_data_size = 'SELECT TABLE_schema AS "Database",SUM(data_length+index_length)/1024 AS "Size(KB)" FROM information_schema.`TABLES`WHERE table_schema="marineranching_system"'
    result1 = db_query.db_query(sql_data_size)
    print(result1)
    pids = psutil.pids()
    cpu_times = psutil.cpu_times()
    mem = psutil.virtual_memory()
    disk = psutil.disk_usage('/')
    print(len(pids))
    print(cpu_times)
    print(mem)
    print(disk)
    sql_device = 'SELECT* FROM device'
    device_result = db_query.db_query(sql_device)
    print(device_result)
    return render_template('data_center.html', data_size=result1[0]['Size(KB)'], pids=len(pids), cpu_times=cpu_times,
                           mem=mem, disk=disk, device_result=device_result)


@app.route('/user_register', methods=['GET', 'POST'])
def user_register():
    if request.method == 'GET':
        return render_template('user_register.html')
    else:
        user_id = request.form.get('user_id')
        password = request.form.get('password1')
        sql = 'select* from user where user_id=%s'
        result = db_query.db_query(sql, [user_id])
        print(result)
        if len(result) == 0:
            sql = 'insert into user(user_id,password) values(%s,%s)'
            db_query.db_query(sql, [user_id, password])
            success = '注册成功'
            return render_template('user_login.html', success=success)
        else:
            error = '注册失败'
            return render_template('user_register.html', error=error)


@app.route('/export-to-excel', methods=['POST'])
def export_to_excel():
    table_name = request.form['table']
    sql = 'select* from ' + table_name
    data = db_query.db_query(sql)
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = table_name
    worksheet.append(list(data[0].keys()))
    for row in data:
        worksheet.append(list(row.values()))
    file_name = f"{table_name}.xlsx"
    workbook.save(file_name)
    return send_file(file_name, as_attachment=True)


if __name__ == '__main__':
    app.run()
